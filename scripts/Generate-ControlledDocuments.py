import argparse,csv,json,hashlib,subprocess,sys,shutil,os
from pathlib import Path
from datetime import datetime,timezone

ap=argparse.ArgumentParser()
ap.add_argument("--ems-root",required=True)
ap.add_argument("--inventory",required=True)
ap.add_argument("--outdir",required=True)
ap.add_argument("--manifest",required=True)
a=ap.parse_args()

root=Path(a.ems_root).resolve()
outdir=Path(a.outdir).resolve()
outdir.mkdir(parents=True,exist_ok=True)

with Path(a.inventory).open(encoding="utf-8-sig",newline="") as f:
    inventory=list(csv.DictReader(f))
if not inventory:
    raise SystemExit("Source inventory is empty.")

def sha256(p):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""):
            h.update(chunk)
    return h.hexdigest()

# We create DOCX with python-docx, which is already part of the EMS build history.
try:
    from docx import Document
except Exception as e:
    raise SystemExit(f"python-docx is required for DOCX generation: {e}")

libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
records=[]
errors=[]

for r in inventory:
    src=root/r["source_path"]
    safe=r["source_name"].replace(" ","_")
    docx_path=outdir/f"{safe}.docx"
    pdf_path=outdir/f"{safe}.pdf"

    try:
        text=src.read_text(encoding="utf-8",errors="replace")
        doc=Document()
        doc.add_heading(r["source_name"],0)
        for block in text.splitlines():
            if block.startswith("# "):
                doc.add_heading(block[2:].strip(),level=1)
            elif block.startswith("## "):
                doc.add_heading(block[3:].strip(),level=2)
            elif block.startswith("### "):
                doc.add_heading(block[4:].strip(),level=3)
            else:
                doc.add_paragraph(block)
        doc.save(docx_path)
    except Exception as e:
        errors.append(f"{r['source_path']}: DOCX generation failed: {e}")
        continue

    pdf_ok=False
    pdf_detail=""
    if libreoffice:
        cp=subprocess.run(
            [libreoffice,"--headless","--convert-to","pdf","--outdir",str(outdir),str(docx_path)],
            cwd=str(root),capture_output=True,text=True
        )
        # LibreOffice names the output after the docx stem.
        pdf_ok=(cp.returncode==0 and pdf_path.exists() and pdf_path.stat().st_size>0)
        pdf_detail=(cp.stdout+"\n"+cp.stderr)[-4000:]
    else:
        pdf_detail="LibreOffice/soffice not found on PATH."

    rec={
        "source_path":r["source_path"],
        "source_sha256":r["sha256"],
        "docx_path":str(docx_path.relative_to(root)).replace("\\","/"),
        "docx_sha256":sha256(docx_path) if docx_path.exists() and docx_path.stat().st_size>0 else "",
        "pdf_path":str(pdf_path.relative_to(root)).replace("\\","/"),
        "pdf_sha256":sha256(pdf_path) if pdf_ok else "",
        "generator":"Wave2C1a Generate-ControlledDocuments.py",
        "generator_version":"2C1a-1.0",
        "generated_at_utc":datetime.now(timezone.utc).isoformat(),
        "docx_size_bytes":docx_path.stat().st_size if docx_path.exists() else 0,
        "pdf_size_bytes":pdf_path.stat().st_size if pdf_path.exists() else 0,
        "pdf_conversion_detail":pdf_detail
    }
    records.append(rec)

for rec in records:
    if not rec["docx_sha256"]:
        errors.append(f"{rec['source_path']}: DOCX output missing/empty.")
    if not rec["pdf_sha256"]:
        errors.append(f"{rec['source_path']}: PDF output missing/empty.")

manifest={
    "wave":"2C.1a",
    "control_id":"EMS-CTRL-072",
    "generator_version":"2C1a-1.0",
    "source_count":len(inventory),
    "record_count":len(records),
    "status":"PASS" if not errors and len(records)==len(inventory) else "FAIL",
    "errors":errors,
    "records":records
}
Path(a.manifest).write_text(json.dumps(manifest,indent=2),encoding="utf-8")
print(json.dumps({
    "source_count":len(inventory),
    "record_count":len(records),
    "docx_count":sum(1 for x in records if x["docx_sha256"]),
    "pdf_count":sum(1 for x in records if x["pdf_sha256"]),
    "status":manifest["status"],
    "errors":errors
},indent=2))
raise SystemExit(0 if manifest["status"]=="PASS" else 2)
